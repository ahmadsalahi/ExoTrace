"""
lib/excel_generator.py — Publication-Grade Styled Excel (.xlsx) Generator for ExoTrace.

Produces beautifully styled, RTL-formatted Excel workbooks with colored headers,
gold/amber comet candidate highlighting, zebra striping, and auto-fitted columns.
"""

import io
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st


@st.cache_data(ttl=120)
def generate_styled_excel(df_records: list[dict], report_title: str, scope_desc: str, lang: str = None) -> bytes:
    """
    Build a professionally styled Excel workbook from records.
    Supports English (LTR) and Arabic (RTL) dynamically.
    """
    if lang is None:
        try:
            from lib.i18n import get_lang
            lang = get_lang()
        except Exception:
            lang = "ar"

    is_en = (lang == "en")

    if not df_records:
        df = pd.DataFrame()
    else:
        df = pd.DataFrame(df_records)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comets_and_Signals" if is_en else "المذنبات_والإشارات"

    # RTL/LTR Orientation and visible gridlines
    ws.views.sheetView[0].rightToLeft = False if is_en else True
    ws.views.sheetView[0].showGridLines = True

    num_cols = max(len(df.columns), 1)

    # 1. Main Title Banner (Row 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    title_prefix = "🔭 ExoTrace Spectroscopic Observatory — " if is_en else "🔭 ExoTrace للأرصاد الفلكية — "
    title_cell = ws.cell(row=1, column=1, value=f"{title_prefix}{report_title}")
    title_cell.font = Font(name="Segoe UI", size=13.5, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0B1120", end_color="0B1120", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. Metadata Subtitle (Row 2)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    if is_en:
        meta_text = (
            f"📅 Export Date: {time.strftime('%Y-%m-%d %H:%M')}  •  "
            f"🎯 Scope: {scope_desc}  •  "
            f"📊 Total Records: {len(df)} signals"
        )
    else:
        meta_text = (
            f"📅 تاريخ الاستخراج: {time.strftime('%Y-%m-%d %H:%M')}  •  "
            f"🎯 النطاق: {scope_desc}  •  "
            f"📊 إجمالي السجلات: {len(df)} إشارة رصدية"
        )
    meta_cell = ws.cell(row=2, column=1, value=meta_text)
    meta_cell.font = Font(name="Segoe UI", size=9.5, bold=False, color="94A3B8")
    meta_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    meta_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    # 3. Spacing (Row 3)
    ws.row_dimensions[3].height = 10

    # Cell Borders
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    header_border = Border(
        left=Side(style="thin", color="0284C7"),
        right=Side(style="thin", color="0284C7"),
        top=Side(style="thin", color="0284C7"),
        bottom=Side(style="medium", color="0369A1"),
    )

    # 4. Table Column Headers (Row 4)
    ws.row_dimensions[4].height = 28
    for col_idx, col_name in enumerate(df.columns, start=1):
        c = ws.cell(row=4, column=col_idx, value=col_name)
        c.font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = header_border

    # Fills & Fonts for data rows
    comet_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    comet_font = Font(name="Segoe UI", size=9.5, bold=True, color="78350F")

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=9.5, color="334155")

    # 5. Data Rows (Row 5+)
    start_row = 5
    comet_count = 0
    for r_idx, (_, row) in enumerate(df.iterrows(), start=start_row):
        ws.row_dimensions[r_idx].height = 22
        cls_txt = str(row.get("التصنيف الرصدي", row.get("Classification", "")))
        is_comet = "مذنّب" in cls_txt or "مرشح" in cls_txt or "Comet" in cls_txt or "candidate" in cls_txt.lower()
        if is_comet:
            comet_count += 1

        for c_idx, col_name in enumerate(df.columns, start=1):
            val = row[col_name]
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border

            if is_comet:
                cell.fill = comet_fill
                cell.font = comet_font
            else:
                cell.fill = alt_fill if (r_idx % 2 == 0) else white_fill
                cell.font = regular_font

            if col_name in ["اسم النجم", "التصنيف الرصدي"]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name in ["Star Name", "Classification", "Observation Summary"]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 6. Summary Footer Row
    summary_row_idx = start_row + len(df) + 1
    ws.row_dimensions[summary_row_idx].height = 24
    ws.merge_cells(start_row=summary_row_idx, start_column=1, end_row=summary_row_idx, end_column=num_cols)
    if is_en:
        footer_text = f"☄️ Total Exocomet Candidates identified in this table: {comet_count}  •  📊 Total Photometric Signals: {len(df)}"
    else:
        footer_text = f"☄️ إجمالي المذنبات المرشحة المكتشفة في هذا الجدول: {comet_count} مذنّب  •  📊 إجمالي الإشارات: {len(df)}"
    footer_cell = ws.cell(row=summary_row_idx, column=1, value=footer_text)
    footer_cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="0369A1")
    footer_cell.fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    footer_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 7. Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < 4 or cell.row == summary_row_idx:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
